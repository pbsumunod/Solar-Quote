"""Validate what clients receive, not just whether an archive was produced."""
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from logic.reports import generate_atomic, generate_csv, generate_excel, generate_pdf, generate_pptx


@pytest.fixture
def proposal():
    rows = [SimpleNamespace(year=year, electricity_rate=14.5, grid_cost_without_solar=20_000,
                            grid_cost_with_solar=8_000, electricity_savings=12_000, maintenance=1_000,
                            replacement=2_000 if year == 10 else 0, net_savings=11_000,
                            cumulative_cash_flow=year * 11_000 - 100_000,
                            discounted_net_savings=10_000, cumulative_discounted_cash_flow=year * 10_000 - 100_000,
                            energy_retention=0.95) for year in range(1, 51)]
    return {
        "quote_id": "SQ-TEST-01", "created_with_version": "1.1.0", "company_name": "Test Solar",
        "client": {"client_name": "Jane Doe", "project_location": "Manila", "project_date": "2026-10-06",
                   "monthly_kwh": 400, "daytime_pct": 70, "nighttime_pct": 30, "electricity_rate": 14.5},
        "consumption": {"daily_kwh": 13.3, "peak_kwh": 2.2, "hourly_kwh": [0.55] * 24},
        "solar": {"total_kw": 4.5, "peak_sun_hours": 4, "efficiency_pct": 80},
        "battery": {"total_kwh": 10, "usable_kwh": 8, "dod_pct": 80, "efficiency_pct": 90},
        "system_design": {"rated_power_kw": 5, "efficiency_pct": 97,
                          "warnings": [(False, "PV array exceeds verified input capacity.")]},
        "simulation": {"scenario_results": {}, "weather_factors": {"Sunny": 100, "Cloudy": 55, "Rainy": 30}},
        "assumptions": {"roi": {"weather_mix": {"Sunny": 60, "Cloudy": 25, "Rainy": 15},
                                "export_rate": 6, "maintenance_method": "percentage", "maintenance_pct": 1,
                                "discount_rate_pct": 5, "pv_degradation_pct": 0.5}},
        "roi": {"monthly_savings": 1000, "annual_savings": 12000, "system_cost": 100000,
                "payback_years": 9.1, "net_benefit": 50000, "roi_pct": 50, "npv": 42000,
                "projection_rows": rows},
    }


def test_pdf_preserves_unicode_and_long_client_fields(tmp_path, proposal):
    proposal["client"]["client_name"] = "王明 – Muñoz ₱ Việt"
    proposal["client"]["project_location"] = "Barangay San Isidro, Pampanga; " * 15
    proposal["company_name"] = "太阳能 – Solar ₱"
    destination = tmp_path / "unicode.pdf"
    generate_pdf(destination, proposal)
    text = "\n".join(page.extract_text() for page in PdfReader(destination).pages)
    assert "王明 – Muñoz ₱ Việt" in text
    assert "太阳能 – Solar ₱" in text
    assert "San Isidro" in text
    assert "SQ-TEST-01" in text


def test_every_automatic_pdf_page_has_watermark(tmp_path, proposal):
    destination = tmp_path / "watermarked.pdf"
    generate_pdf(destination, proposal, watermark="SolarQuote Free")
    pages = PdfReader(destination).pages
    assert len(pages) >= 4
    assert all("SolarQuote Free" in page.extract_text() for page in pages)


def test_all_formats_contain_equipment_assumptions_and_warnings(tmp_path, proposal):
    csv_path, excel_path, pdf_path, ppt_path = [tmp_path / name for name in ("quote.csv", "quote.xlsx", "quote.pdf", "quote.pptx")]
    generate_csv(csv_path, proposal)
    generate_excel(excel_path, proposal)
    generate_pdf(pdf_path, proposal)
    generate_pptx(ppt_path, proposal)
    workbook = load_workbook(excel_path)
    excel_text = "\n".join(str(cell.value) for sheet in workbook for row in sheet for cell in row)
    pdf_text = "\n".join(page.extract_text() for page in PdfReader(pdf_path).pages)
    ppt_text = "\n".join(shape.text for slide in Presentation(ppt_path).slides for shape in slide.shapes if shape.has_text_frame)
    for text in (csv_path.read_text(encoding="utf-8-sig"), excel_text, pdf_text, ppt_text):
        assert "SQ-TEST-01" in text
        assert "Solar array (kWp)" in text
        assert "Battery total (kWh)" in text
        assert "Inverter rated power (kW)" in text
        assert "PV array exceeds verified input capacity." in text
        assert "Export Rate" in text
        assert "Maintenance Method" in text
        assert "Pv Degradation Pct" in text
        assert "matched to inverter rating" not in text
    assert workbook["Selected equipment"]["B2"].value == 4.5
    assert "Replacement (PHP)" in [cell.value for cell in workbook["Financial projection"][1]]


def test_pptx_works_outside_repository_and_without_optional_photo(tmp_path, proposal, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr("logic.reports.ROI_SLIDE_PHOTO", tmp_path / "not-installed.jpg")
    generate_pptx(tmp_path / "portable.pptx", proposal)
    assert len(Presentation(tmp_path / "portable.pptx").slides) > 5


def test_atomic_generation_cleans_partial_and_preserves_existing_file(tmp_path, proposal, monkeypatch):
    destination = tmp_path / "quote.pdf"
    destination.write_bytes(b"existing proposal")
    def fail(path, *_args, **_kwargs):
        path = str(path)
        with open(path, "wb") as file:
            file.write(b"incomplete")
        raise OSError("out of storage")
    monkeypatch.setattr("logic.reports.generate_pdf", fail)
    with pytest.raises(OSError, match="out of storage"):
        generate_atomic("pdf", destination, proposal)
    assert destination.read_bytes() == b"existing proposal"
    assert list(tmp_path.iterdir()) == [destination]


def test_spreadsheet_all_user_text_is_escaped(tmp_path, proposal):
    proposal["client"]["project_location"] = "  =HYPERLINK(\"https://example.test\")"
    proposal["roi"]["net_benefit"] = -5000
    generate_excel(tmp_path / "safe.xlsx", proposal)
    workbook = load_workbook(tmp_path / "safe.xlsx")
    values = dict(workbook["Client profile"].values)
    assert values["Project Location"].startswith("'")
    generate_csv(tmp_path / "safe.csv", proposal)
    assert "Net Benefit,-5000" in (tmp_path / "safe.csv").read_text(encoding="utf-8-sig")


def test_financial_sensitivity_is_readable_in_all_formats(tmp_path, proposal):
    proposal["roi"]["sensitivity_model"] = "Approximate yield scaling; lifetime assumptions held constant."
    proposal["roi"]["sensitivity_rows"] = [{"scenario": "Solar yield -20%", "npv": 12345,
                                            "payback_years": 12.3, "roi_pct": 30}]
    for suffix, generator in (("csv", generate_csv), ("xlsx", generate_excel), ("pdf", generate_pdf), ("pptx", generate_pptx)):
        destination = tmp_path / f"sensitivity.{suffix}"
        generator(destination, proposal)
        if suffix == "csv":
            text = destination.read_text(encoding="utf-8-sig")
        elif suffix == "xlsx":
            text = "\n".join(str(cell.value) for row in load_workbook(destination)["Financial sensitivity"] for cell in row)
        elif suffix == "pdf":
            text = "\n".join(page.extract_text() for page in PdfReader(destination).pages)
        else:
            text = "\n".join(shape.text for slide in Presentation(destination).slides for shape in slide.shapes if shape.has_text_frame)
        assert "Solar yield -20%" in text
        assert "12,345.00" in text
        assert "12.3 years" in text
        assert "Approximate yield scaling" in text
