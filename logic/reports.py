"""Portable proposal exports from one immutable, auditable quote snapshot."""

import csv
import hashlib
import json
import os
from datetime import datetime
from pathlib import Path

from logic.calculations import HOUR_LABELS

DISCLAIMER = (
    "Solar system sizing and financial returns are estimates based on the stated assumptions. "
    "Weather, shading, actual demand, equipment specifications, degradation, installation, "
    "and utility requirements affect results. Hourly energy does not establish an instantaneous "
    "peak or motor starting load. Final design and utility/export eligibility require verification "
    "by a qualified solar/electrical professional."
)
SCENARIO_ORDER = ["Sunny", "Cloudy", "Rainy"]
REPORT_SCHEMA_VERSION = "2"
ASSET_ROOT = Path(__file__).resolve().parent.parent / "assets"
FONT_PATH = ASSET_ROOT / "fonts" / "NotoSans-Regular.ttf"
ROI_SLIDE_PHOTO = ASSET_ROOT / "images" / "roi_slide.jpg"
PRIMARY_RGB = (46, 126, 52)
TEXT_RGB = (28, 28, 24)
MUTED_RGB = (72, 71, 62)
BAND_RGB = (243, 241, 234)

PROJECTION_COLUMNS = [
    ("Year", "year"), ("Rate (PHP/kWh)", "electricity_rate"),
    ("Grid cost without solar (PHP)", "grid_cost_without_solar"),
    ("Grid cost with solar (PHP)", "grid_cost_with_solar"),
    ("Electricity savings (PHP)", "electricity_savings"),
    ("Maintenance (PHP)", "maintenance"), ("Replacement (PHP)", "replacement"),
    ("Net savings (PHP)", "net_savings"), ("Cumulative cash flow (PHP)", "cumulative_cash_flow"),
    ("Discounted net savings (PHP)", "discounted_net_savings"),
    ("Cumulative discounted cash flow (PHP)", "cumulative_discounted_cash_flow"),
    ("Energy retention", "energy_retention"),
]


def _spreadsheet_text(value):
    """Prevent client supplied strings from becoming spreadsheet formulas."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _display(value):
    if value is None:
        return "N/A"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def _flatten(values, prefix=""):
    for key, value in values.items():
        label = f"{prefix} / {key.replace('_', ' ').title()}" if prefix else key.replace("_", " ").title()
        if isinstance(value, dict):
            yield from _flatten(value, label)
        else:
            yield label, value


def proposal_metadata(data):
    identity = {key: data.get(key, {}) for key in ("client", "solar", "battery", "system_design", "assumptions")}
    digest = hashlib.sha256(json.dumps(identity, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()[:12]
    return [
        ("Quote ID", data.get("quote_id") or f"SQ-{digest.upper()}"),
        ("Created with version", data.get("created_with_version", "1.0.0")),
        ("Report schema", REPORT_SCHEMA_VERSION),
        ("Generated at", data.get("generated_at") or datetime.now().astimezone().isoformat(timespec="seconds")),
    ]


def report_sections(data):
    """The same equipment, risks, inputs and financial summary in every format."""
    client, solar, battery, inverter, roi = (data[key] for key in ("client", "solar", "battery", "system_design", "roi"))
    equipment = [
        ("Solar array (kWp)", solar["total_kw"]),
        ("Battery total (kWh)", battery["total_kwh"]),
        ("Battery usable (kWh)", battery["usable_kwh"]),
        ("Inverter rated power (kW)", inverter["rated_power_kw"]),
    ]
    for section, values in (("Solar", solar), ("Battery", battery), ("Inverter", inverter)):
        equipment.extend(_flatten({key: value for key, value in values.items()
                                   if key not in {"warnings", "equipment", "total_kw", "total_kwh", "usable_kwh", "rated_power_kw"}}, section))
    design = data.get("equipment_design") or inverter.get("equipment") or solar.get("equipment") or battery.get("equipment")
    if design:
        equipment.extend(_flatten(design, "Design"))
    assumptions = data.get("assumptions") or {
        "client": client,
        "solar": {key: value for key, value in solar.items() if key != "total_kw"},
        "battery": {key: value for key, value in battery.items() if key not in {"total_kwh", "usable_kwh"}},
        "inverter": {key: value for key, value in inverter.items() if key != "warnings"},
        "simulation": {key: value for key, value in data["simulation"].items() if key != "scenario_results"},
        "financial": {key: value for key, value in roi.items() if key not in {"projection_rows", "sensitivity_rows"}},
    }
    notes = [("OK" if ok else "CHECK REQUIRED", text) for ok, text in inverter.get("warnings", [])]
    if not notes:
        notes = [("Review required", "No equipment verification was supplied. Confirm manufacturer specifications and site conditions.")]
    financial = [(key.replace("_", " ").title(), "Not within modeled horizon" if value is None and "payback" in key else value) for key, value in roi.items()
                 if key not in {"projection_rows", "weather_mix", "sensitivity_rows", "sensitivity_model"}]
    scenarios = []
    for name in SCENARIO_ORDER:
        result = data["simulation"]["scenario_results"].get(name)
        if result is not None:
            scenarios.extend([(f"{name} / AC load (kWh/day)", result.total_consumption_kwh),
                              (f"{name} / DC solar harvest (kWh/day)", result.total_solar_kwh),
                              (f"{name} / Grid purchase (kWh/day)", result.total_grid_purchase_kwh),
                              (f"{name} / Grid export (kWh/day)", result.total_grid_export_kwh),
                              (f"{name} / Curtailed solar (kWh/day)", result.total_curtailed_kwh),
                              (f"{name} / Conversion loss (kWh/day)", result.total_conversion_loss_kwh)])
    sections = [
        ("Proposal identification", proposal_metadata(data)),
        ("Client profile", list(_flatten(client))),
        ("Selected equipment", equipment),
        ("Engineering notes", notes),
        ("Consumption summary", list(_flatten({key: value for key, value in data["consumption"].items() if key != "hourly_kwh"}))),
        ("Weather scenario totals", scenarios),
        ("Assumptions", list(_flatten(assumptions))),
        ("Savings & ROI", financial),
    ]
    sensitivity = roi.get("sensitivity_rows", data.get("sensitivity_rows", []))
    if sensitivity:
        pairs = [("Model", roi.get("sensitivity_model", data.get("sensitivity_model", "Sensitivity estimates; other assumptions held constant.")))]
        for case in sensitivity:
            npv = case.get("npv")
            payback = case.get("payback_years")
            npv_text = f"PHP {npv:,.2f}" if isinstance(npv, (int, float)) else "N/A"
            payback_text = f"{payback:.1f} years" if isinstance(payback, (int, float)) else "Not within modeled horizon"
            roi_value = case.get("roi_pct")
            roi_text = f"{roi_value:.1f}%" if isinstance(roi_value, (int, float)) else "N/A"
            pairs.append((case.get("scenario", "Scenario"), f"NPV: {npv_text}; Payback: {payback_text}; ROI: {roi_text}"))
        sections.append(("Financial sensitivity", pairs))
    return sections


def _projection_values(row):
    return [getattr(row, attr, None) for _title, attr in PROJECTION_COLUMNS]


def generate_excel(path, data):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for title, pairs in report_sections(data):
        sheet = wb.create_sheet(title[:31])
        sheet.append(["Item", "Value"])
        for label, value in pairs:
            sheet.append([label, _spreadsheet_text(_display(value)) if not isinstance(value, (int, float)) else value])
        sheet.column_dimensions["A"].width = 45
        sheet.column_dimensions["B"].width = 65
    sheet = wb.create_sheet("Hourly consumption")
    sheet.append(["Hour", "Consumption (kWh)"])
    for label, value in zip(HOUR_LABELS, data["consumption"].get("hourly_kwh", [])):
        sheet.append([label, value])
    sheet = wb.create_sheet("Weather scenarios")
    sheet.append(["Scenario", "AC load (kWh/day)", "DC solar (kWh/day)", "Grid purchase (kWh/day)",
                  "Grid export (kWh/day)", "Curtailed solar (kWh/day)", "Conversion loss (kWh/day)"])
    for name in SCENARIO_ORDER:
        result = data["simulation"]["scenario_results"].get(name)
        if result is None:
            continue
        sheet.append([name, result.total_consumption_kwh, result.total_solar_kwh,
                      result.total_grid_purchase_kwh, result.total_grid_export_kwh,
                      result.total_curtailed_kwh, result.total_conversion_loss_kwh])
        hourly = wb.create_sheet(f"Hourly simulation {name}")
        fields = ["consumption_kwh", "solar_kwh", "solar_to_load_kwh", "battery_charge_kwh", "battery_discharge_kwh",
                  "grid_purchase_kwh", "grid_export_kwh", "curtailed_kwh", "soc_end_kwh", "soc_end_pct"]
        hourly.append(["Hour"] + fields)
        for hour in result.hours:
            hourly.append([hour.hour_label] + [getattr(hour, field) for field in fields])
    sheet = wb.create_sheet("Financial projection")
    sheet.append([title for title, _attr in PROJECTION_COLUMNS])
    for row in data["roi"]["projection_rows"]:
        sheet.append(_projection_values(row))
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    wb.save(path)


def generate_csv(path, data):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        # Legacy financial-only callers remain valid; the app always supplies a full snapshot.
        sections = report_sections(data) if all(key in data for key in ("solar", "battery", "system_design", "simulation", "consumption")) else [
            ("Client", [("Client Name", data["client"]["client_name"]), ("Monthly Consumption (kWh)", data["client"]["monthly_kwh"])]),
            ("Savings & ROI", [(key.replace("_", " ").title(), value) for key, value in data["roi"].items() if key != "projection_rows"]),
        ]
        for title, pairs in sections:
            writer.writerow([title])
            writer.writerows([[label, value if isinstance(value, (int, float)) else _spreadsheet_text(_display(value))]
                              for label, value in pairs])
            writer.writerow([])
        writer.writerow([title for title, _attr in PROJECTION_COLUMNS])
        writer.writerows(_projection_values(row) for row in data["roi"]["projection_rows"])
        writer.writerow([])
        writer.writerow(["Disclaimer", DISCLAIMER])


def _pdf_watermark(pdf, text):
    saved_x, saved_y = pdf.x, pdf.y
    with pdf.rotation(45, x=pdf.w / 2, y=pdf.h / 2):
        pdf.set_font(pdf.font_family or "Helvetica", "B", 48)
        pdf.set_text_color(215, 215, 215)
        pdf.set_xy(0, pdf.h / 2)
        pdf.cell(pdf.w, 0, text, align="C")
    pdf.set_xy(saved_x, saved_y)


def _pdf_heading(pdf, text):
    pdf.set_font("SolarQuote", "B", 14)
    pdf.set_text_color(*PRIMARY_RGB)
    pdf.multi_cell(pdf.epw, 8, text, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)


def _pdf_kv_table(pdf, pairs):
    from fpdf.fonts import FontFace
    pdf.set_font("SolarQuote", "", 9)
    pdf.set_text_color(*TEXT_RGB)
    with pdf.table(col_widths=(0.36, 0.64), first_row_as_headings=False,
                   line_height=5, borders_layout="HORIZONTAL_LINES") as table:
        for index, (label, value) in enumerate(pairs):
            text = _display(value)
            # A single unbounded user-field row cannot span pages in fpdf2 tables.
            # Continue large values in bounded rows rather than clipping/failing.
            chunks = [text[start:start + 600] for start in range(0, len(text), 600)] or [""]
            for part, chunk in enumerate(chunks):
                row = table.row(style=FontFace(fill_color=BAND_RGB) if index % 2 else None)
                row.cell(str(label) + (" (continued)" if part else ""))
                row.cell(chunk)
    pdf.ln(4)


def generate_pdf(path, data, roi_chart_image=None, scenario_images=None, watermark=None):
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    class ProposalPDF(FPDF):
        def header(self):
            if watermark:
                _pdf_watermark(self, watermark)

        def footer(self):
            self.set_y(-11)
            self.set_font("SolarQuote", "", 8)
            self.set_text_color(*MUTED_RGB)
            self.cell(0, 5, f"{dict(proposal_metadata(data))['Quote ID']} | Page {self.page_no()}", align="C")

    pdf = ProposalPDF()
    # Static OFL font, Latin extended + punctuation + peso + CJK. No module monkeypatches.
    for style in ("", "B", "I"):
        font = ASSET_ROOT / "fonts" / "NotoSans-Bold.ttf" if style == "B" else FONT_PATH
        pdf.add_font("SolarQuote", style=style, fname=str(font))
    pdf.add_font("NotoSC", fname=str(ASSET_ROOT / "fonts" / "NotoSansSC.ttf"))
    pdf.set_fallback_fonts(["NotoSC"], exact_match=False)
    pdf.set_font("SolarQuote", "", 10)
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    pdf.set_font("SolarQuote", "B", 20)
    pdf.set_text_color(*PRIMARY_RGB)
    pdf.multi_cell(pdf.epw, 10, data.get("company_name") or "Solar Proposal", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("SolarQuote", "", 12)
    pdf.multi_cell(pdf.epw, 7, f"Solar System Proposal for {data['client']['client_name'] or 'Prospective Client'}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    for title, pairs in report_sections(data):
        _pdf_heading(pdf, title)
        _pdf_kv_table(pdf, pairs)
    for name in SCENARIO_ORDER:
        image = (scenario_images or {}).get(name)
        if image:
            pdf.add_page()
            _pdf_heading(pdf, f"{name} day simulation")
            pdf.image(image, w=pdf.epw)
            pdf.ln(4)
            pdf.set_font("SolarQuote", "", 9)
            pdf.multi_cell(pdf.epw, 5, "Hourly energy (kWh), 6 AM to 5 AM: red AC load; yellow DC solar harvest; blue AC grid (+ export, - import); dashed green stored battery energy. DOD floor marks minimum stored energy.", new_x="LMARGIN", new_y="NEXT")
    if roi_chart_image:
        pdf.add_page()
        _pdf_heading(pdf, "Cumulative cash flow (PHP); horizontal axis: years")
        pdf.image(roi_chart_image, w=pdf.epw)
    pdf.add_page()
    _pdf_heading(pdf, "Financial projection (PHP)")
    # Keep financial columns readable in two tables; each repeats year for joining.
    for columns in (PROJECTION_COLUMNS[:9], [PROJECTION_COLUMNS[0]] + PROJECTION_COLUMNS[9:]):
        pdf.set_font("SolarQuote", "", 7)
        with pdf.table(line_height=5, text_align="LEFT", headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=PRIMARY_RGB)) as table:
            header = table.row()
            for title, _attr in columns:
                header.cell(title)
            for projection in data["roi"]["projection_rows"]:
                row = table.row()
                for _title, attr in columns:
                    value = getattr(projection, attr, None)
                    row.cell(f"{value:,.2f}" if isinstance(value, float) else ("-" if value is None else str(value)))
        pdf.ln(5)
    pdf.set_font("SolarQuote", "I", 8)
    pdf.multi_cell(pdf.epw, 4, DISCLAIMER, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def _ppt_text(slide, left, top, width, height, text, size=16, bold=False, color="1C1C18"):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for index, line in enumerate(str(text).split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Noto Sans SC"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
    return box


def generate_pptx(path, data, chart_images=None, roi_chart_image=None):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def slide(title):
        page = prs.slides.add_slide(prs.slide_layouts[6])
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = RGBColor.from_string("F3F1EA")
        _ppt_text(page, 0.6, 0.35, 12.1, 0.7, title, 28, True, "2E7E34")
        _ppt_text(page, 0.6, 7.05, 12.1, 0.3, f"{dict(proposal_metadata(data))['Quote ID']} | SolarQuote | {len(prs.slides)}", 10)
        return page

    cover = slide(data.get("company_name") or "Solar Proposal")
    _ppt_text(cover, 0.8, 2.0, 11.7, 3.2, f"Solar System Proposal for {data['client']['client_name']}\n{data['client'].get('project_date', '')}", 24, True)
    for title, pairs in report_sections(data):
        # Bound text length and line budget so long addresses/notes continue instead of clipping.
        import textwrap
        lines = []
        for label, value in pairs:
            lines.extend(textwrap.wrap(f"{label}: {_display(value)}", width=95, break_long_words=True) or [""])
        for start in range(0, len(lines), 14):
            page = slide(title + (" (continued)" if start else ""))
            _ppt_text(page, 0.7, 1.2, 12, 5.65, "\n".join(lines[start:start + 14]), 15)
    for name in SCENARIO_ORDER:
        image = (chart_images or {}).get(name)
        if image:
            page = slide(f"{name} day simulation")
            page.shapes.add_picture(str(image), Inches(0.6), Inches(1.15), width=Inches(12.1))
            _ppt_text(page, 0.6, 6.5, 12.1, 0.5, "6 AM to 5 AM | Red: AC load | Yellow: DC solar | Blue: AC grid (+ export, - import) | Green: stored battery. kWh.", 12)
    from pptx.util import Pt
    # Split the wide annual projection into related groups with a year in each.
    for columns in (PROJECTION_COLUMNS[:6], [PROJECTION_COLUMNS[0]] + PROJECTION_COLUMNS[6:9],
                    [PROJECTION_COLUMNS[0]] + PROJECTION_COLUMNS[9:]):
        rows = data["roi"]["projection_rows"]
        for start in range(0, len(rows), 10):
            group = rows[start:start + 10]
            page = slide("Financial projection (PHP)")
            table = page.shapes.add_table(len(group) + 1, len(columns), Inches(0.6), Inches(1.15), Inches(12.1), Inches(5.6)).table
            for col, (title, _attr) in enumerate(columns):
                table.cell(0, col).text = title
            for index, row in enumerate(group, 1):
                for col, (_title, attr) in enumerate(columns):
                    value = getattr(row, attr, None)
                    table.cell(index, col).text = f"{value:,.2f}" if isinstance(value, float) else ("-" if value is None else str(value))
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.name = "Noto Sans SC"
                        paragraph.font.size = Pt(11)
    if roi_chart_image:
        page = slide("Cumulative cash flow (PHP); horizontal axis: years")
        page.shapes.add_picture(str(roi_chart_image), Inches(0.6), Inches(1.15), width=Inches(12.1))
    if ROI_SLIDE_PHOTO.is_file():
        page = slide("Savings summary")
        page.shapes.add_picture(str(ROI_SLIDE_PHOTO), Inches(7.5), Inches(1.4), width=Inches(5))
        roi = data["roi"]
        payback = f"{roi['payback_years']:.1f} years" if roi.get("payback_years") is not None else "Not within modeled horizon"
        _ppt_text(page, 0.7, 1.7, 6.5, 4.5, f"System cost: PHP {roi['system_cost']:,.2f}\nNet benefit: PHP {roi['net_benefit']:,.2f}\nPayback: {payback}", 20)
    page = slide("Review before issuing this proposal")
    _ppt_text(page, 0.8, 1.5, 11.7, 4.8, DISCLAIMER, 20)
    prs.save(str(path))


def generate_atomic(kind, path, data, *, roi_chart_image=None, scenario_images=None, watermark=None):
    """Never leave a partial proposal at the final filename on an export failure."""
    import tempfile
    descriptor, temp_path = tempfile.mkstemp(prefix=".solarquote-", suffix=Path(path).suffix, dir=str(Path(path).parent))
    os.close(descriptor)
    try:
        if kind == "pdf":
            generate_pdf(temp_path, data, roi_chart_image, scenario_images, watermark)
        elif kind == "pptx":
            generate_pptx(temp_path, data, scenario_images, roi_chart_image)
        elif kind == "excel":
            generate_excel(temp_path, data)
        elif kind == "csv":
            generate_csv(temp_path, data)
        else:
            raise ValueError("Unsupported export format")
        os.replace(temp_path, path)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
